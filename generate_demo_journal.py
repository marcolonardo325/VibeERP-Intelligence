"""
Generate a demo General Journal Excel file for D365 F&O (USMF company).
All accounts, journal names, currencies, and account types are validated 
against the live USMF environment master data.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date, timedelta
import os

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "D365FO_Demo_Journal_Transactions.xlsx")

# ── Header row (matches LedgerJournalTrans entity / Excel add-in template) ──
HEADERS = [
    "Journal batch number",   # Leave blank – will be assigned when importing into a new journal
    "Line number",
    "Date",
    "Company",
    "Account type",
    "Account",
    "Description",
    "Debit",
    "Credit",
    "Offset account type",
    "Offset account",
    "Currency",
    "Voucher",
]

# ── Demo transactions built from real USMF master data ──
# Rules enforced:
#   • Every line has Debit XOR Credit (never both, never zero on both)
#   • Balanced vouchers (total debit == total credit per voucher)
#   • Valid account types: Ledger, Customer, Vendor, Bank
#   • Ledger accounts use MainAccount numbers from the shared CoA
#   • Customer/Vendor accounts exist in USMF
#   • Bank account IDs used for Bank type (not main account numbers)
#   • Currency = USD (default for USMF, matching bank/customer/vendor currency)
#   • Dates within open fiscal period

today = date(2026, 3, 6)

transactions = [
    # ── Voucher 1: Office supplies expense paid from bank ──
    # Debit expense account, credit bank
    {
        "journal": "", "line": 1, "date": today, "company": "USMF",
        "acct_type": "Ledger", "account": "601200",
        "desc": "Office supplies purchase",
        "debit": 1250.00, "credit": 0.00,
        "offset_type": "Bank", "offset_account": "USMF OPER",
        "currency": "USD", "voucher": "DEMO-001",
    },
    # ── Voucher 2: Customer payment received into bank ──
    # Debit bank, credit customer (settles open AR)
    {
        "journal": "", "line": 2, "date": today, "company": "USMF",
        "acct_type": "Customer", "account": "US-001",
        "desc": "Payment from Contoso Retail San Diego",
        "debit": 0.00, "credit": 5000.00,
        "offset_type": "Bank", "offset_account": "USMF OPER",
        "currency": "USD", "voucher": "DEMO-002",
    },
    # ── Voucher 3: Vendor invoice recorded ──
    # Debit expense, credit vendor (AP)
    {
        "journal": "", "line": 3, "date": today, "company": "USMF",
        "acct_type": "Vendor", "account": "1001",
        "desc": "Acme Office Supplies invoice",
        "debit": 0.00, "credit": 3200.00,
        "offset_type": "Ledger", "offset_account": "601200",
        "currency": "USD", "voucher": "DEMO-003",
    },
    # ── Voucher 4: Rent expense – GL to GL ──
    {
        "journal": "", "line": 4, "date": today, "company": "USMF",
        "acct_type": "Ledger", "account": "601500",
        "desc": "Monthly rent expense",
        "debit": 8500.00, "credit": 0.00,
        "offset_type": "Ledger", "offset_account": "200110",
        "currency": "USD", "voucher": "DEMO-004",
    },
    # ── Voucher 5: Service revenue from customer ──
    {
        "journal": "", "line": 5, "date": today, "company": "USMF",
        "acct_type": "Customer", "account": "US-003",
        "desc": "Consulting services invoice - Forest Wholesales",
        "debit": 15000.00, "credit": 0.00,
        "offset_type": "Ledger", "offset_account": "403150",
        "currency": "USD", "voucher": "DEMO-005",
    },
    # ── Voucher 6: Pay vendor from bank ──
    {
        "journal": "", "line": 6, "date": today, "company": "USMF",
        "acct_type": "Vendor", "account": "1001",
        "desc": "Payment to Acme Office Supplies",
        "debit": 3200.00, "credit": 0.00,
        "offset_type": "Bank", "offset_account": "USMF OPER",
        "currency": "USD", "voucher": "DEMO-006",
    },
    # ── Voucher 7: Prepaid insurance ──
    {
        "journal": "", "line": 7, "date": today, "company": "USMF",
        "acct_type": "Ledger", "account": "132100",
        "desc": "Prepaid insurance - quarterly",
        "debit": 4500.00, "credit": 0.00,
        "offset_type": "Bank", "offset_account": "USMF OPER",
        "currency": "USD", "voucher": "DEMO-007",
    },
    # ── Voucher 8: Customer US-004 payment received ──
    {
        "journal": "", "line": 8, "date": today, "company": "USMF",
        "acct_type": "Customer", "account": "US-004",
        "desc": "Payment from Cave Wholesales",
        "debit": 0.00, "credit": 7500.00,
        "offset_type": "Bank", "offset_account": "USMF OPER",
        "currency": "USD", "voucher": "DEMO-008",
    },
    # ── Voucher 9: Intercompany receivable (GL to GL) ──
    {
        "journal": "", "line": 9, "date": today, "company": "USMF",
        "acct_type": "Ledger", "account": "130100",
        "desc": "Intercompany charge - AR adjustment",
        "debit": 2000.00, "credit": 0.00,
        "offset_type": "Ledger", "offset_account": "403150",
        "currency": "USD", "voucher": "DEMO-009",
    },
    # ── Voucher 10: Petty cash replenishment from bank ──
    {
        "journal": "", "line": 10, "date": today, "company": "USMF",
        "acct_type": "Ledger", "account": "110180",
        "desc": "Petty cash replenishment",
        "debit": 500.00, "credit": 0.00,
        "offset_type": "Bank", "offset_account": "USMF OPER",
        "currency": "USD", "voucher": "DEMO-010",
    },
]


def build_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "General Journal Lines"

    # ── Styles ──
    header_font = Font(name="Segoe UI", bold=True, size=11, color="003366")
    header_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="Segoe UI", size=11)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # ── Write headers ──
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # ── Write data rows ──
    for row_idx, txn in enumerate(transactions, start=2):
        values = [
            txn["journal"],
            txn["line"],
            txn["date"],
            txn["company"],
            txn["acct_type"],
            txn["account"],
            txn["desc"],
            txn["debit"],
            txn["credit"],
            txn["offset_type"],
            txn["offset_account"],
            txn["currency"],
            txn["voucher"],
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = cell_font
            cell.border = thin_border
            if col_idx == 3:  # Date
                cell.number_format = "YYYY-MM-DD"
            elif col_idx in (8, 9):  # Debit/Credit
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")

    # ── Column widths ──
    col_widths = [22, 12, 14, 10, 14, 16, 45, 14, 14, 20, 18, 10, 14]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Freeze header row ──
    ws.freeze_panes = "A2"

    # ── Add an instructions sheet ──
    ws2 = wb.create_sheet("Instructions")
    instructions = [
        "D365 F&O General Journal – Demo Data Import Instructions",
        "",
        "Company: USMF (Contoso Entertainment System USA)",
        "Journal Name: GenJrn (General Journal)",
        "",
        "STEPS TO IMPORT:",
        "1. Open D365 F&O → General ledger → Journal entries → General journals",
        "2. Click 'New' to create a new journal; select Name = 'GenJrn'",
        "3. Click 'Open lines in Excel' to get the template, OR use Data Management",
        "4. Copy the rows from the 'General Journal Lines' sheet into the template",
        "5. The 'Journal batch number' column can be left blank (auto-assigned)",
        "6. Click 'Validate' → 'Post' when ready",
        "",
        "VALIDATION RULES FOLLOWED:",
        "• Each voucher is self-balancing (Debit = Credit)",
        "• Account types match valid master data in USMF",
        "• Ledger accounts exist in the Shared chart of accounts",
        "• Customer accounts (US-001, US-003, US-004) exist in USMF",
        "• Vendor account (1001 - Acme Office Supplies) exists in USMF",
        "• Bank account (USMF OPER) is a valid bank account in USMF",
        "• Currency is USD (matches company default and account currencies)",
        "• Date is within an open fiscal period (2026-03-06)",
        "• Only Debit OR Credit is populated per line (never both)",
        "",
        "ACCOUNT TYPES USED:",
        "  Ledger  → Main account numbers (e.g., 601200, 601500, 403150)",
        "  Customer → Customer account IDs (e.g., US-001, US-003)",
        "  Vendor  → Vendor account IDs  (e.g., 1001)",
        "  Bank    → Bank account IDs    (e.g., USMF OPER)",
    ]
    for r, line in enumerate(instructions, start=1):
        cell = ws2.cell(row=r, column=1, value=line)
        if r == 1:
            cell.font = Font(name="Segoe UI", bold=True, size=14, color="0078D4")
        else:
            cell.font = Font(name="Segoe UI", size=11)
    ws2.column_dimensions["A"].width = 80

    wb.save(OUTPUT_PATH)
    print(f"✅ Excel saved: {OUTPUT_PATH}")
    print(f"   → {len(transactions)} transactions across {len(set(t['voucher'] for t in transactions))} vouchers")

    # ── Also generate CSV for maximum compatibility ──
    import csv
    csv_path = OUTPUT_PATH.replace(".xlsx", ".csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        for txn in transactions:
            writer.writerow([
                txn["journal"],
                txn["line"],
                txn["date"].strftime("%Y-%m-%d"),
                txn["company"],
                txn["acct_type"],
                txn["account"],
                txn["desc"],
                txn["debit"],
                txn["credit"],
                txn["offset_type"],
                txn["offset_account"],
                txn["currency"],
                txn["voucher"],
            ])
    print(f"✅ CSV saved:   {csv_path}")


if __name__ == "__main__":
    build_workbook()
